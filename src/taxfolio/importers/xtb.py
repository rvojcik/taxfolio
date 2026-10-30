"""Parser for XTB (xStation 5) cash operations: CSV or the Excel account statement.

Supported layouts (all read by column name, so extra or reordered columns are fine):

* legacy CSV ``ID;Type;Time;Symbol;Comment;Amount`` with ``DD.MM.YYYY HH:MM:SS`` times,
* current export ``Type;Ticker;Instrument;Time;Amount;ID;Comment;Product`` with a few
  metadata rows (account number, currency, period) above the header and a ``Total`` row;
  ``;`` with decimal comma or ``,`` with decimal point,
* the same table on the "Cash Operations" sheet of the Excel statement (.xlsx).

Stock and ETF trades are read from the comment (``OPEN BUY 34/42.5658 @ 11.7480`` = 34 pieces of
an order for 42.5658 at 11.748 in the instrument's currency). ``Amount`` is always in the account
currency and already includes the currency conversion fee, so trades are stored in the account
currency with price = amount / quantity. The comment format does not depend on the language of
the export; the type column is used only for rows without a trade or dividend comment.

The export carries no ISIN, so instruments are keyed by the XTB symbol (``AAPL.US``) and the
dividend source country is estimated from the exchange suffix. CFD and forex rows (profit/loss,
swap, rollover) are not supported and are reported as warnings.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, replace
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from taxfolio.domain.models import (
    AccountInfo,
    AssetClass,
    CashTransaction,
    CashType,
    Instrument,
    ParsedFile,
    Trade,
)
from taxfolio.i18n import _

FORMAT = "xtb_cash_operations"
BROKER = "XTB"

_TRADE = re.compile(r"\b(OPEN|CLOSE)\s+(BUY|SELL)\s+([\d.,]+)(?:\s*/\s*[\d.,]+)?\s*@\s*([\d.,]+)", re.I)
_DIVIDEND = re.compile(r"\b[A-Z]{3}\s+[\d.,]+\s*/\s*SHR\b", re.I)
_WHT = re.compile(r"\bWHT\b", re.I)
_TIME_FORMATS = (
    "%d.%m.%Y %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%d/%m/%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%Y-%m-%d %H:%M",
)

# Exchange suffix of an XTB symbol → (country of the exchange, trading currency).
# The country is only a guess for the dividend source: an ETF listed in Germany is usually Irish.
EXCHANGES: dict[str, tuple[str, str]] = {
    "US": ("US", "USD"),
    "UK": ("GB", "GBP"),
    "DE": ("DE", "EUR"),
    "FR": ("FR", "EUR"),
    "NL": ("NL", "EUR"),
    "ES": ("ES", "EUR"),
    "IT": ("IT", "EUR"),
    "BE": ("BE", "EUR"),
    "PT": ("PT", "EUR"),
    "FI": ("FI", "EUR"),
    "IE": ("IE", "EUR"),
    "PL": ("PL", "PLN"),
    "CZ": ("CZ", "CZK"),
    "DK": ("DK", "DKK"),
    "SE": ("SE", "SEK"),
    "NO": ("NO", "NOK"),
    "CH": ("CH", "CHF"),
}
CURRENCIES = {"EUR", "USD", "CZK", "PLN", "GBP", "HUF", "RON", "CHF"}

# Header names (lower case) → field
_COLUMNS = {
    "id": "id",
    "type": "type",
    "time": "time",
    "symbol": "symbol",
    "ticker": "symbol",
    "instrument": "instrument",
    "comment": "comment",
    "amount": "amount",
    "product": "product",
}
_REQUIRED = {"type", "time", "amount", "comment"}


@dataclass(frozen=True)
class XtbRow:
    line_no: int
    id: str
    type: str
    time: datetime
    symbol: str
    instrument: str
    comment: str
    amount: Decimal


# --------------------------------------------------------------------------- detection


def _header_fields(cells: list[str]) -> dict[str, int] | None:
    found = {}
    for i, cell in enumerate(cells):
        name = _COLUMNS.get(str(cell or "").strip().lower())
        if name and name not in found:
            found[name] = i
    return found if _REQUIRED <= found.keys() and "symbol" in found else None


def is_xtb_csv(head: str) -> bool:
    for line in head.splitlines()[:20]:
        for delimiter in (";", ",", "\t"):
            if _header_fields(line.split(delimiter)):
                return True
    return False


def is_xtb_xlsx(path: Path) -> bool:
    return path.suffix.lower() == ".xlsx" and _cash_sheet(path) is not None


# --------------------------------------------------------------------------- reading


def _cash_sheet(path: Path) -> list[list[object]] | None:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:  # not a readable workbook
        return None
    try:
        for ws in wb.worksheets:
            name = ws.title.lower()
            if "cash" in name and "operation" in name:
                return [list(r) for r in ws.iter_rows(values_only=True)]
        return None
    finally:
        wb.close()


def _csv_rows(path: Path) -> list[list[object]]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = "\n".join(text.splitlines()[:30])
    delimiter = max(
        (";", ",", "\t"), key=lambda d: max((len(line.split(d)) for line in sample.splitlines()), default=0)
    )
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]


def parse_number(value: object) -> Decimal:
    """``-500,00``, ``-500.00``, ``1 234,56``, ``1,234.56`` or a spreadsheet number."""
    if isinstance(value, (int, float, Decimal)):
        # Spreadsheet floats carry binary noise (23.050000000000001)
        return Decimal(str(value)).quantize(Decimal("0.00000001")).normalize()
    s = str(value or "").strip().replace(" ", "").replace(" ", "")
    if not s:
        return Decimal(0)
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation as exc:
        raise ValueError(f"not a number: {value!r}") from exc


def parse_time(value: object) -> datetime:
    if isinstance(value, datetime):
        return value.replace(microsecond=0, tzinfo=None)
    s = str(value or "").strip()
    for fmt in _TIME_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    raise ValueError(f"unknown time format: {value!r}")


def _cell(row: list[object], index: int | None) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _metadata(rows: list[list[object]], label: str) -> str:
    """Value next to (or below) a metadata label such as 'Account number' or 'Currency'."""
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            if str(cell or "").strip().lower().startswith(label):
                right = next((str(v).strip() for v in row[c + 1 :] if str(v or "").strip()), "")
                if right:
                    return right
                if r + 1 < len(rows) and c < len(rows[r + 1]) and str(rows[r + 1][c] or "").strip():
                    return str(rows[r + 1][c]).strip()
    return ""


def _read_table(
    rows: list[list[object]], name: str, warnings: list[str]
) -> tuple[list[XtbRow], list[list[object]]]:
    """Data rows and the metadata rows above the header; unreadable rows become warnings."""
    header = next(
        ((h, f) for h, row in enumerate(rows) if (f := _header_fields([str(c or "") for c in row]))), None
    )
    if header is None:
        raise ValueError(_("{file}: v souboru chybí tabulka peněžních operací XTB.").format(file=name))
    h, fields = header

    out = []
    for offset, row in enumerate(rows[h + 1 :], start=h + 2):
        kind = _cell(row, fields["type"])
        if not kind or kind.lower() == "total" or any(str(c or "").strip().lower() == "total" for c in row):
            continue  # blank line or the closing sum
        try:
            time, amount = parse_time(row[fields["time"]]), parse_number(row[fields["amount"]])
        except (ValueError, IndexError) as exc:
            warnings.append(
                _("XTB řádek {line}: nelze přečíst ({error}) – přeskočeno.").format(line=offset, error=exc)
            )
            continue
        out.append(
            XtbRow(
                line_no=offset,
                id=_cell(row, fields.get("id")),
                type=kind,
                time=time,
                symbol=_cell(row, fields.get("symbol")),
                instrument=_cell(row, fields.get("instrument")),
                comment=_cell(row, fields["comment"]),
                amount=amount,
            )
        )
    return out, rows[:h]


# --------------------------------------------------------------------------- interpretation


def exchange_of(symbol: str) -> tuple[str, str] | None:
    suffix = symbol.rsplit(".", 1)[-1].upper() if "." in symbol else ""
    return EXCHANGES.get(suffix)


def is_cfd(symbol: str) -> bool:
    """Index and commodity CFDs: ``US30.cash``, ``DE40``, ``GOLD``, ``EURUSD`` (no exchange suffix)."""
    if not symbol:
        return False
    suffix = symbol.rsplit(".", 1)[-1] if "." in symbol else ""
    return not suffix or not suffix.isupper()


def _infer_currency(rows: list[XtbRow]) -> str | None:
    """Account currency from trades priced in their exchange currency: there amount = qty × price."""
    votes: Counter[str] = Counter()
    for r in rows:
        m = _TRADE.search(r.comment)
        ex = exchange_of(r.symbol)
        if not m or not ex or not r.amount:
            continue
        value = parse_number(m.group(3)) * parse_number(m.group(4))
        if abs(value - abs(r.amount)) <= max(Decimal("0.011"), abs(r.amount) * Decimal("0.0005")):
            votes[ex[1]] += 1
    return votes.most_common(1)[0][0] if votes else None


def _account(path: Path, meta: list[list[object]], rows: list[XtbRow], warnings: list[str]) -> AccountInfo:
    number = _metadata(meta, "account")
    if not re.fullmatch(r"\d+", number or ""):
        m = re.search(r"(?:account|konto|ucet|účet)?_?(?:ike_|ikze_)?(\d{5,})", path.stem, re.I)
        number = m.group(1) if m else ""
    currency = _metadata(meta, "currency").upper()
    if currency not in CURRENCIES:
        tokens = {t.upper() for t in re.split(r"[^A-Za-z]+", path.stem)}
        currency = next(iter(tokens & CURRENCIES), "") or (_infer_currency(rows) or "")
    if not currency:
        raise ValueError(
            _("{file}: nelze zjistit měnu účtu XTB. Přidejte ji do názvu souboru, např. „{example}“.").format(
                file=path.name, example=f"EUR_{path.name}"
            )
        )
    if not number:
        warnings.append(
            _("{file}: číslo účtu XTB nenalezeno, data se uloží pod účet „{account}“.").format(
                file=path.name, account=f"{BROKER}-{currency}"
            )
        )
        number = currency
    return AccountInfo(f"{BROKER}-{number}", f"{BROKER} {number}", currency)


def parse(path: Path) -> ParsedFile:
    raw = _cash_sheet(path) if path.suffix.lower() == ".xlsx" else _csv_rows(path)
    if raw is None:
        raise ValueError(_("{file}: v souboru chybí tabulka peněžních operací XTB.").format(file=path.name))
    warnings: list[str] = []
    rows, meta = _read_table(raw, path.name, warnings)
    account = _account(path, meta, rows, warnings)
    return interpret(rows, account, warnings)


def interpret(rows: list[XtbRow], account: AccountInfo, warnings: list[str] | None = None) -> ParsedFile:
    acc, ccy = account.account_id, account.base_currency
    result = ParsedFile(source_format=FORMAT, account=account)
    names: dict[str, str] = {}
    trades: list[Trade] = []
    fees: list[XtbRow] = []
    skipped: Counter[str] = Counter()
    unknown: list[XtbRow] = []

    def cash(r: XtbRow, kind: CashType, symbol: str | None = None, action_id: str | None = None) -> None:
        result.cash.append(
            CashTransaction(
                account_id=acc,
                type=kind,
                currency=ccy,
                date=r.time.date(),
                amount=r.amount,
                description=r.comment or r.type,
                symbol=symbol,
                conid=f"{BROKER}:{symbol}" if symbol else None,
                action_id=action_id,
                transaction_id=r.id or None,
            )
        )

    seen: set[tuple[str, str, str, Decimal]] = set()
    for r in sorted(rows, key=lambda r: (r.time, r.id)):
        kind = r.type.lower()
        if r.id:
            # Overlapping exports pasted together repeat operations; XTB ids are unique.
            ident = (r.id, r.type, r.comment, r.amount)
            if ident in seen:
                continue
            seen.add(ident)
        if is_cfd(r.symbol):
            skipped[r.type] += 1
            continue
        if r.symbol and r.instrument:
            names.setdefault(r.symbol, r.instrument)
        if m := _TRADE.search(r.comment):
            action, side = m.group(1).upper(), m.group(2).upper()
            quantity = parse_number(m.group(3))
            if side != "BUY" or not quantity or not r.symbol:
                unknown.append(r)
                continue
            signed = quantity if action == "OPEN" else -quantity
            trades.append(
                Trade(
                    account_id=acc,
                    asset_class=AssetClass.STOCK,
                    symbol=r.symbol,
                    currency=ccy,
                    trade_datetime=r.time,
                    quantity=signed,
                    price=(abs(r.amount) / quantity).quantize(Decimal("0.000001")),
                    proceeds=r.amount,
                    commission=Decimal(0),
                    commission_currency=ccy,
                    codes=("O",) if action == "OPEN" else ("C",),
                    conid=f"{BROKER}:{r.symbol}",
                    trade_id=r.id or None,
                    transaction_id=r.id or None,
                )
            )
        elif "spin" in kind:
            unknown.append(r)
        elif _WHT.search(r.comment) or ("withholding" in kind and r.symbol):
            cash(r, CashType.WITHHOLDING_TAX, r.symbol, f"{BROKER}:{acc}:{r.symbol}:{r.time.isoformat()}")
        elif _DIVIDEND.search(r.comment) or "divid" in kind:
            cash(r, CashType.DIVIDEND, r.symbol, f"{BROKER}:{acc}:{r.symbol}:{r.time.isoformat()}")
        elif "interest" in kind or "interest" in r.comment.lower():
            taxed = "tax" in kind or "tax" in r.comment.lower()
            cash(r, CashType.WITHHOLDING_TAX if taxed else CashType.INTEREST)
        elif "deposit" in kind:
            cash(r, CashType.DEPOSIT)
        elif "withdraw" in kind:
            cash(r, CashType.WITHDRAWAL)
        elif "transfer" in kind:
            cash(r, CashType.OTHER)
        elif any(word in kind for word in ("commission", "fee", "iftt", "ftt")):
            fees.append(r)
        elif any(word in kind for word in ("profit/loss", "swap", "rollover", "close trade", "cfd")):
            skipped[r.type] += 1
        else:
            unknown.append(r)

    # Commissions and transaction taxes belong to the trade of the same symbol and second.
    by_moment: dict[tuple[str, datetime], list[int]] = defaultdict(list)
    for i, t in enumerate(trades):
        by_moment[(t.symbol, t.trade_datetime)].append(i)
    for f in fees:
        targets = by_moment.get((f.symbol, f.time))
        if targets:
            i = targets[0]
            trades[i] = replace(trades[i], commission=trades[i].commission + f.amount)
        else:
            cash(f, CashType.FEE, f.symbol or None)
    result.trades = trades

    symbols = sorted({t.symbol for t in trades} | {c.symbol for c in result.cash if c.symbol})
    for symbol in symbols:
        ex = exchange_of(symbol)
        result.instruments.append(
            Instrument(
                symbol=symbol,
                asset_class=AssetClass.STOCK,
                description=names.get(symbol, ""),
                conid=f"{BROKER}:{symbol}",
                # "XTB:US" – the rules list these as regulated markets (SK time test)
                listing_exchange=f"{BROKER}:{symbol.rsplit('.', 1)[-1]}" if "." in symbol else None,
                issuer_country=ex[0] if ex else None,
            )
        )

    out = result.warnings
    for kind, count in sorted(skipped.items()):
        out.append(
            _("XTB: {count}× „{type}“ (CFD/forex) – tyto obchody aplikace nepodporuje, přeskočeno.").format(
                count=count, type=kind
            )
        )
    for r in unknown:
        out.append(
            _(
                "XTB řádek {line}: nerozpoznaná operace „{type}“ {symbol} {comment} {amount} – přeskočeno."
            ).format(line=r.line_no, type=r.type, symbol=r.symbol, comment=r.comment, amount=r.amount)
        )
    dividend_symbols = sorted({c.symbol for c in result.cash if c.type == CashType.DIVIDEND and c.symbol})
    guessed = [s for s in dividend_symbols if (exchange_of(s) or ("", ""))[0] != "US"]
    if guessed:
        out.append(
            _(
                "XTB neuvádí ISIN: stát zdroje dividend {symbols} je odhadnut podle burzy. "
                "U ETF (často irské) jej ověřte."
            ).format(symbols=", ".join(guessed))
        )
    if warnings is not None:
        out[:0] = warnings
    dates: list[date] = [t.trade_datetime.date() for t in trades] + [c.date for c in result.cash]
    if dates:
        result.period_start, result.period_end = min(dates), max(dates)
    return result
