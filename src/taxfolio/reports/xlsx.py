"""Render a report Document to an Excel workbook (one sheet per table + summary sheet)."""

from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from taxfolio.i18n import _
from taxfolio.reports.document import Document

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(title: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("-", title)[:31] or "List"
    name, n = base, 2
    while name in used:
        suffix = f" ({n})"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(name)
    return name


def write_xlsx(doc: Document, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    summary = _("Souhrn")
    used = {summary}
    ws.title = summary
    ws.append([doc.title])
    ws["A1"].font = Font(bold=True, size=14)
    if doc.subtitle:
        ws.append([doc.subtitle])
    ws.append([_("Vygenerováno {when}").format(when=f"{doc.generated_at:%d.%m.%Y %H:%M}")])
    ws.append([])
    for paragraph in doc.summary:
        ws.append([paragraph])
    if doc.notes:
        ws.append([])
        ws.append([_("Upozornění")])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for note in doc.notes:
            ws.append([note])
    ws.column_dimensions["A"].width = 140

    for table in doc.tables:
        sheet = wb.create_sheet(_sheet_name(table.title, used))
        sheet.append(table.headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in table.rows:
            # Excel has no decimal type; floats are fine for display, values were rounded upstream
            sheet.append([float(v) if isinstance(v, Decimal) else v for v in row])
        for idx, header in enumerate(table.headers, start=1):
            width = max([len(header), *(len(str(r[idx - 1] or "")) for r in table.rows)], default=10)
            sheet.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 60)
            for (cell,) in sheet.iter_rows(min_row=2, min_col=idx, max_col=idx):
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
                elif hasattr(cell.value, "year"):
                    cell.number_format = "DD.MM.YYYY"
        sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
