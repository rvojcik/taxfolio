from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import Engine

from taxfolio.fx.converter import CzkMethod
from taxfolio.fx.sources import CNB, ECB
from taxfolio.reports.export import export
from taxfolio.reports.tax_documents import cz_document, sk_document
from taxfolio.services import import_service, rates_service
from taxfolio.services.tax_cz import build_cz_report
from taxfolio.services.tax_sk import build_sk_report
from tests.test_realized import _flat_rates


@pytest.fixture
def loaded(engine: Engine, fixtures: Path) -> Engine:
    import_service.import_file(engine, fixtures / "flex_2025.xml")
    rates_service.ensure_rates(
        engine,
        [2025, 2026],
        fetch_cnb=_flat_rates(CNB, "USD", "22"),
        fetch_ecb_years=_flat_rates(ECB, "USD", "1.1"),
    )
    return engine


def test_cz_export_xlsx_and_pdf(loaded: Engine, tmp_path: Path) -> None:
    report, calendar = build_cz_report(loaded, 2025, CzkMethod.CNB_DAILY)
    doc = cz_document(report, calendar, "U0000001")
    export(doc, tmp_path / "cz.xlsx")
    export(doc, tmp_path / "cz.pdf")
    wb = openpyxl.load_workbook(tmp_path / "cz.xlsx")
    assert wb.sheetnames[:3] == ["Souhrn", "Do přiznání", "Prodeje cenných papírů"]
    sales = wb["Prodeje cenných papírů"]
    assert sales.cell(2, 2).value == "AAPL" and sales.max_row == 2
    assert (tmp_path / "cz.pdf").read_bytes().startswith(b"%PDF")


def test_sk_export_and_unknown_format(loaded: Engine, tmp_path: Path) -> None:
    report, calendar = build_sk_report(loaded, 2025)
    doc = sk_document(report, calendar)
    export(doc, tmp_path / "sk.xlsx")
    assert "Dividendy – jednotlivé výplaty" in openpyxl.load_workbook(tmp_path / "sk.xlsx").sheetnames
    with pytest.raises(ValueError):
        export(doc, tmp_path / "sk.docx")
